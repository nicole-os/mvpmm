"""
docx_exporter.py
Generates a .docx campaign package from pitchplanner campaign data.

Actual campaign data structure (from pitcher.py / app.js):
  waves.wave_1  → { target_data: {...}, exclusive_offer, contingency, angle_note }
  waves.wave_2  → [ { target_data: {...}, angle_note }, ... ]
  waves.wave_3  → [ { target_data: {...}, angle_note, format_suggestion }, ... ]

  target_data → {
    publication, fit_score, tier, beat, audience,
    pitch: { subject_line, body, send_date, embargoed_briefing,
             key_talking_points, personalization_notes,
             companion_content_recommended, exclusive_offer_line, follow_on_hook },
    recent_headlines, known_authors, audience_hook
  }
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
from datetime import datetime


# ── Colour palette ────────────────────────────────────────────
PLUM      = RGBColor(0x4A, 0x34, 0x53)
GREY_MED  = RGBColor(0x88, 0x88, 0x88)
BODY      = RGBColor(0x22, 0x22, 0x22)


# ── Helpers ───────────────────────────────────────────────────

def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = PLUM
    run.font.name = 'Calibri'
    run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 11)


def _add_label(doc: Document, label: str, value, indent: float = 0):
    if not value:
        return
    if isinstance(value, bool):
        value = "Yes" if value else "No"
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = BODY
    r2 = p.add_run(str(value))
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def _add_body(doc: Document, text: str, indent: float = 0, italic: bool = False):
    if not text:
        return
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    for run in p.runs:
        run.font.size = Pt(9.5)
        run.font.color.rgb = BODY
        if italic:
            run.italic = True


def _add_divider(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'C8BFD0')
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_pitch_block(doc: Document, target_data: dict, wave_label: str, angle_note: str = ''):
    """Render one target + its pitch email."""
    if not target_data:
        return

    pub = target_data.get("publication", "Unknown")
    pitch = target_data.get("pitch") or {}
    score = target_data.get("fit_score", "")
    tier = target_data.get("tier", "")
    send_date = pitch.get("send_date", "")
    subject = pitch.get("subject_line", "")
    body = pitch.get("body", "")
    embargoed = pitch.get("embargoed_briefing", False)
    talking_points = pitch.get("key_talking_points") or []
    personalization = pitch.get("personalization_notes", "")
    companion = pitch.get("companion_content_recommended", "")
    exclusive_line = pitch.get("exclusive_offer_line", "")
    follow_on = pitch.get("follow_on_hook", "")
    audience_hook = target_data.get("audience_hook", "")
    authors = target_data.get("known_authors") or []

    # Pub name + wave label
    p_pub = doc.add_paragraph()
    p_pub.paragraph_format.space_before = Pt(10)
    p_pub.paragraph_format.space_after = Pt(2)
    r_pub = p_pub.add_run(pub)
    r_pub.bold = True
    r_pub.font.size = Pt(12)
    r_pub.font.color.rgb = PLUM

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(6)
    r_meta = p_meta.add_run(f"{wave_label}  ·  {send_date}  ·  Tier {tier}  ·  Score {score}/10")
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = GREY_MED

    if audience_hook:
        _add_label(doc, "Audience hook", audience_hook)
    if angle_note:
        _add_label(doc, "Angle", angle_note)
    if authors:
        _add_label(doc, "Known contributors", ", ".join(authors))

    # Pitch email
    if subject or body:
        p_email_head = doc.add_paragraph()
        p_email_head.paragraph_format.space_before = Pt(8)
        p_email_head.paragraph_format.space_after = Pt(2)
        p_email_head.paragraph_format.left_indent = Inches(0.2)
        r_eh = p_email_head.add_run("Draft Pitch Email")
        r_eh.bold = True
        r_eh.font.size = Pt(9.5)
        r_eh.font.color.rgb = PLUM

        if subject:
            _add_label(doc, "Subject", subject, indent=0.2)
        if body:
            _add_body(doc, body, indent=0.2)

        _add_label(doc, "Embargoed briefing offered", "Yes" if embargoed else "No", indent=0.2)

        if talking_points:
            p_tp_head = doc.add_paragraph()
            p_tp_head.paragraph_format.left_indent = Inches(0.2)
            r_tp = p_tp_head.add_run("Key talking points:")
            r_tp.bold = True
            r_tp.font.size = Pt(9)
            for pt in talking_points:
                p_tp = doc.add_paragraph(f"• {pt}")
                p_tp.paragraph_format.left_indent = Inches(0.4)
                for run in p_tp.runs:
                    run.font.size = Pt(9)

        if personalization:
            _add_label(doc, "Personalization rationale", personalization, indent=0.2)
        if companion:
            _add_label(doc, "Attach / offer", companion, indent=0.2)
        if exclusive_line:
            _add_label(doc, "Exclusive offer line", exclusive_line, indent=0.2)
        if follow_on:
            _add_label(doc, "Follow-on hook", follow_on, indent=0.2)

    _add_divider(doc)


def build_campaign_docx(campaign: dict, session: dict) -> bytes:
    """Build and return a .docx campaign document as bytes."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    r_title = title_p.add_run("pitchplanner — Campaign Export")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = PLUM

    date_p = doc.add_paragraph()
    r_date = date_p.add_run(f"Generated {datetime.now().strftime('%B %d, %Y')}")
    r_date.font.size = Pt(10)
    r_date.font.color.rgb = GREY_MED
    doc.add_paragraph()

    # ── News Assessment ──────────────────────────────────────
    news = campaign.get("news_analysis") or {}
    if news:
        _add_heading(doc, "News Assessment", level=1)
        _add_label(doc, "Core story", news.get("core_story"))
        _add_label(doc, "Why now", news.get("why_now"))
        _add_label(doc, "Who cares", news.get("who_cares"))
        _add_label(doc, "Newsworthiness", f"{news.get('newsworthiness_score', '')}/10")
        _add_label(doc, "Recommended approach", news.get("recommended_approach"))
        angles = news.get("story_angles") or []
        if angles:
            p_h = doc.add_paragraph()
            p_h.add_run("Story angles:").bold = True
            p_h.runs[0].font.size = Pt(9.5)
            for angle in angles:
                text = angle.get("angle") or angle.get("title") or str(angle) if isinstance(angle, dict) else str(angle)
                p_a = doc.add_paragraph(f"• {text}")
                p_a.paragraph_format.left_indent = Inches(0.2)
                for run in p_a.runs:
                    run.font.size = Pt(9.5)
        _add_divider(doc)

    # ── Campaign Plan ────────────────────────────────────────
    plan = campaign.get("campaign_plan") or {}
    if plan:
        _add_heading(doc, "Campaign Plan", level=1)
        _add_label(doc, "Strategy", plan.get("campaign_summary") or plan.get("overall_strategy"))
        _add_label(doc, "Key message", plan.get("key_message"))
        _add_divider(doc)

    # ── Waves ────────────────────────────────────────────────
    waves = campaign.get("waves") or {}

    # Wave 1
    w1_entry = waves.get("wave_1")
    if w1_entry:
        _add_heading(doc, "Wave 1 — Exclusive", level=1)
        w1_plan = plan.get("wave_1") or {}
        if w1_plan.get("timing_label"):
            _add_body(doc, w1_plan["timing_label"], italic=True)
        exclusive_offer = w1_entry.get("exclusive_offer")
        if exclusive_offer:
            _add_label(doc, "What you're offering exclusively", exclusive_offer)
        _add_pitch_block(doc, w1_entry.get("target_data"), "Wave 1 · Exclusive", w1_entry.get("angle_note", ""))

    # Wave 2
    w2_list = waves.get("wave_2") or []
    if w2_list:
        _add_heading(doc, "Wave 2 — Launch Day", level=1)
        w2_plan = plan.get("wave_2") or {}
        if w2_plan.get("timing_label"):
            _add_body(doc, w2_plan["timing_label"], italic=True)
        if w2_plan.get("wave_2_note"):
            _add_body(doc, w2_plan["wave_2_note"])
        for entry in w2_list:
            _add_pitch_block(doc, entry.get("target_data"), "Wave 2 · Launch", entry.get("angle_note", ""))

    # Wave 3
    w3_list = waves.get("wave_3") or []
    if w3_list:
        _add_heading(doc, "Wave 3 — Follow-on", level=1)
        w3_plan = plan.get("wave_3") or {}
        if w3_plan.get("timing_label"):
            _add_body(doc, w3_plan["timing_label"], italic=True)
        if w3_plan.get("wave_3_strategy"):
            _add_body(doc, w3_plan["wave_3_strategy"])
        for entry in w3_list:
            _add_pitch_block(doc, entry.get("target_data"), "Wave 3 · Follow-on", entry.get("angle_note", ""))

    # ── Press Release ────────────────────────────────────────
    pr_data = campaign.get("press_release") or {}
    if pr_data:
        _add_heading(doc, "Press Release", level=1)
        pr_text = pr_data.get("press_release", "")
        if pr_text:
            p = doc.add_paragraph()
            r = p.add_run(pr_text)
            r.font.size = Pt(9.5)
            r.font.color.rgb = BODY
            p.paragraph_format.space_after = Pt(8)

        memo = pr_data.get("pr_firm_brief", "")
        if memo:
            _add_heading(doc, "PR Firm Brief", level=2)
            _add_body(doc, memo)

        embargo = pr_data.get("embargo_management_protocol", "")
        if embargo:
            _add_heading(doc, "Embargo Management Protocol", level=2)
            _add_body(doc, embargo)

    # Footer
    doc.add_paragraph()
    footer_p = doc.add_paragraph("Generated by pitchplanner — AI-powered media pitch generator")
    footer_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer_p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GREY_MED

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
